"""GET-INJS — import d'un fichier externe d'inscrits (cohortes hors parcours).

**Besoin métier.** Le commanditaire reçoit, pour une année académique, un
fichier (CSV ou Excel) listant les étudiants inscrits — parfois produit par un
service qui n'a pas utilisé la chaîne de parcours de l'application. Il faut
pouvoir générer les emplois du temps de ces cohortes **sans** exiger une
réinscription interne, et **sans** fabriquer de dossiers étudiants fictifs.

**Ce que fait ce service.**

1. il lit le fichier (``.csv`` séparateur détecté, ``.xlsx``/``.xlsm`` via
   ``openpyxl`` déjà présent au projet — aucune dépendance nouvelle) ;
2. il reconnaît les colonnes par **synonymes tolérants** (accents, casse,
   espaces, variantes françaises) : un fichier métier n'a jamais l'en-tête
   exact attendu ;
3. il résout chaque ligne vers la chaîne LMD **réelle** (année, formation,
   parcours, niveau, semestre, groupe, vague) ; les entités manquantes sont
   créées **uniquement** pour les formations du périmètre INJS ;
4. il écrit des ``InscriptionImportee`` — jamais d'``Etudiant`` ni
   d'``Inscription`` : la donnée reste identifiée comme importée ;
5. il est **idempotent** : réimporter le même fichier met à jour les lignes
   existantes (clé : année + matricule) au lieu de les dupliquer ;
6. il fonctionne en **analyse** (``appliquer=False``) : le même rapport est
   produit sans aucune écriture, ce qui permet de valider un fichier client
   avant de l'intégrer.

**Ce que ce service ne fait pas.** Il n'invente ni formation, ni niveau, ni
semestre : une ligne dont la formation est introuvable (ou hors périmètre INJS)
est **rejetée avec son motif**, ligne par ligne. Le rapport expose toujours les
comptes réels — jamais un taux embellissant.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction

# ── Motifs de rejet (explicites, affichables tels quels) ─────────────────────
REJET_MATRICULE_MANQUANT = 'MATRICULE_MANQUANT'
REJET_NOM_MANQUANT = 'NOM_MANQUANT'
REJET_FORMATION_ABSENTE = 'FORMATION_ABSENTE'
REJET_FORMATION_HORS_PERIMETRE = 'FORMATION_HORS_PERIMETRE'
REJET_NIVEAU_ABSENT = 'NIVEAU_ABSENT'
REJET_SEMESTRE_ABSENT = 'SEMESTRE_ABSENT'
REJET_FICHIER_VIDE = 'FICHIER_VIDE'
REJET_FORMAT_NON_SUPPORTE = 'FORMAT_NON_SUPPORTE'
REJET_COLONNES_ABSENTES = 'COLONNES_OBLIGATOIRES_ABSENTES'

#: Extensions acceptées (aucune bibliothèque nouvelle : ``openpyxl`` est déjà
#: une dépendance du projet pour l'import Excel des formations).
EXTENSIONS_CSV = ('.csv', '.txt')
EXTENSIONS_XLSX = ('.xlsx', '.xlsm')

#: Synonymes d'en-têtes. La comparaison se fait sur une clé normalisée
#: (sans accent, minuscules, séparateurs ramenés à « _ »).
SYNONYMES = {
    'matricule': (
        'matricule', 'numero_matricule', 'no_matricule', 'n_matricule',
        'num_matricule', 'numero', 'n_matricule_etudiant',
        'matricule_etudiant', 'code_etudiant', 'identifiant',
    ),
    'nom': ('nom', 'nom_etudiant', 'nom_de_famille', 'patronyme'),
    'prenoms': ('prenoms', 'prenom', 'prenom_s', 'prenoms_etudiant'),
    'sexe': ('sexe', 'genre'),
    'date_naissance': ('date_naissance', 'date_de_naissance', 'naissance', 'ddn'),
    'formation': (
        'formation', 'filiere', 'cycle', 'formation_cycle', 'intitule_formation',
        'code_formation', 'ref_formation',
    ),
    'filiere_code': ('filiere_code', 'code_filiere', 'f_code', 'code'),
    'parcours': ('parcours', 'specialite', 'option', 'parcours_specialite'),
    'niveau': ('niveau', 'niveau_etude', 'niveau_d_etude', 'annee_etude'),
    'semestre': ('semestre', 'semestre_etude', 's'),
    'groupe': ('groupe', 'groupe_pedagogique', 'code_groupe', 'groupe_td'),
    'vague': ('vague', 'vague_inscription', 'code_vague'),
    'statut': ('statut', 'statut_inscription', 'etat', 'etat_inscription'),
}

#: Colonnes sans lesquelles aucune ligne n'est exploitable.
COLONNES_OBLIGATOIRES = ('matricule', 'nom', 'formation', 'niveau')

#: Correspondance niveau texte → code du référentiel (S1..S10 → L1..M2).
#: Volontairement limitée au référentiel INJS LMD : rien d'autre n'est deviné.
_SEMESTRE_VERS_NIVEAU = {
    'S1': 'L1', 'S2': 'L1', 'S3': 'L2', 'S4': 'L2', 'S5': 'L3', 'S6': 'L3',
    'S7': 'M1', 'S8': 'M1', 'S9': 'M2', 'S10': 'M2',
}


def normaliser_cle(valeur):
    """Normalise un en-tête ou une valeur de référence en clé comparable."""
    texte = str(valeur or '').strip().lower()
    texte = unicodedata.normalize('NFKD', texte)
    texte = ''.join(c for c in texte if not unicodedata.combining(c))
    texte = re.sub(r'[^a-z0-9]+', '_', texte).strip('_')
    return texte


def _champ(cle):
    """Retrouve le champ canonique correspondant à une clé normalisée."""
    for champ, synonymes in SYNONYMES.items():
        if cle in synonymes:
            return champ
    return None


@dataclass
class LigneErreur:
    ligne: int
    motif: str
    detail: str

    def as_dict(self):
        return {'ligne': self.ligne, 'motif': self.motif, 'detail': self.detail}


@dataclass
class RapportImport:
    """Rapport d'import — toujours rendu, même en cas d'échec global."""

    applique: bool
    fichier: str = ''
    empreinte: str = ''
    lignes_lues: int = 0
    lignes_valides: int = 0
    lignes_rejetees: int = 0
    creees: int = 0
    mises_a_jour: int = 0
    erreurs: list = field(default_factory=list)
    colonnes_reconnues: dict = field(default_factory=dict)
    colonnes_ignorees: list = field(default_factory=list)
    cohortes: list = field(default_factory=list)
    effectif_total: int = 0
    motif_echec: str | None = None
    detail: dict = field(default_factory=dict)

    @property
    def succes(self):
        return self.motif_echec is None and self.lignes_valides > 0

    def as_dict(self):
        return {
            'applique': self.applique,
            'fichier': self.fichier,
            'empreinte': self.empreinte,
            'lignes_lues': self.lignes_lues,
            'lignes_valides': self.lignes_valides,
            'lignes_rejetees': self.lignes_rejetees,
            'creees': self.creees,
            'mises_a_jour': self.mises_a_jour,
            'erreurs': [e.as_dict() for e in self.erreurs[:200]],
            'erreurs_tronquees': max(0, len(self.erreurs) - 200),
            'colonnes_reconnues': self.colonnes_reconnues,
            'colonnes_ignorees': self.colonnes_ignorees,
            'cohortes': self.cohortes,
            'effectif_total': self.effectif_total,
            'motif_echec': self.motif_echec,
            'detail': self.detail,
            'succes': self.succes,
        }


# ── Lecture du fichier ───────────────────────────────────────────────────────

def _lire_csv(contenu_bytes):
    """Lit un CSV en détectant le séparateur (``;`` très fréquent en France)."""
    for encodage in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            texte = contenu_bytes.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 ne lève jamais
        raise ValueError('Encodage non reconnu')

    premiere = texte.split('\n', 1)[0]
    separateur = max((';', ',', '\t'), key=premiere.count)
    lecteur = csv.DictReader(io.StringIO(texte), delimiter=separateur)
    return [dict(ligne) for ligne in lecteur]


def _lire_xlsx(contenu_fichier):
    """Lit la première feuille d'un classeur Excel (en-tête = 1re ligne)."""
    from openpyxl import load_workbook

    classeur = load_workbook(contenu_fichier, read_only=True, data_only=True)
    feuille = classeur.worksheets[0]
    lignes = feuille.iter_rows(values_only=True)
    try:
        entetes = [str(c) if c is not None else '' for c in next(lignes)]
    except StopIteration:
        return []
    resultat = []
    for ligne in lignes:
        if ligne is None or all(c is None for c in ligne):
            continue
        resultat.append({
            entetes[i]: ligne[i] if i < len(ligne) else None
            for i in range(len(entetes))
        })
    classeur.close()
    return resultat


def lire_fichier(nom_fichier, contenu):
    """Retourne la liste des lignes brutes, ou lève ``ValueError`` motivée."""
    nom = (nom_fichier or '').lower()
    if nom.endswith(EXTENSIONS_XLSX):
        flux = io.BytesIO(contenu) if isinstance(contenu, (bytes, bytearray)) else contenu
        return _lire_xlsx(flux)
    if nom.endswith(EXTENSIONS_CSV) or not nom:
        if isinstance(contenu, str):
            contenu = contenu.encode('utf-8')
        return _lire_csv(contenu)
    raise ValueError(REJET_FORMAT_NON_SUPPORTE)


def map_colonnes(entetes):
    """Associe chaque champ canonique à l'en-tête réel du fichier."""
    correspondance = {}
    ignorees = []
    for entete in entetes:
        cle = normaliser_cle(entete)
        champ = _champ(cle)
        if champ and champ not in correspondance:
            correspondance[champ] = entete
        elif cle:
            ignorees.append(entete)
    return correspondance, ignorees


# ── Résolution des référentiels ──────────────────────────────────────────────

def _valeur(ligne, correspondance, champ):
    entete = correspondance.get(champ)
    if entete is None:
        return ''
    brut = ligne.get(entete)
    if brut is None:
        return ''
    return str(brut).strip()


def _trouver_formation(valeur, code_filiere=''):
    """Résout une formation **du périmètre INJS** (jamais un cycle legacy)."""
    from formations.models import RefFormation
    from formations.referentiel_injs_l6 import formations_injs

    perimetre = formations_injs()

    if code_filiere:
        code = str(code_filiere).strip().upper()
        trouvee = perimetre.filter(filiere_code__iexact=code).first()
        if trouvee is not None:
            return trouvee

    texte = str(valeur or '').strip()
    if not texte:
        return None
    trouvee = perimetre.filter(intitule__iexact=texte).first()
    if trouvee is not None:
        return trouvee
    trouvee = perimetre.filter(code__iexact=texte).first()
    if trouvee is not None:
        return trouvee

    # Recherche tolérante : « Education Motricite » ≈ « Éducation Motricité ».
    cible = normaliser_cle(texte)
    for formation in perimetre:
        if normaliser_cle(formation.intitule) == cible:
            return formation
    for formation in perimetre:
        if cible and cible in normaliser_cle(formation.intitule):
            return formation
    return None


def _trouver_niveau(valeur, semestre_texte=''):
    """Résout un niveau (``L1``…) ou le déduit d'un semestre (``S3`` → ``L2``)."""
    from scolarite.models import Niveau

    texte = str(valeur or '').strip()
    if texte:
        niveau = Niveau.objects.filter(code__iexact=texte).first()
        if niveau is not None:
            return niveau
        for niveau in Niveau.objects.all():
            if normaliser_cle(niveau.libelle) == normaliser_cle(texte):
                return niveau
        # « Licence 1 », « LICENCE 1 » → code L1
        nombre = re.search(r'(\d)', texte)
        if nombre:
            cycle = 'M' if 'master' in normaliser_cle(texte) else 'L'
            niveau = Niveau.objects.filter(code__iexact=f'{cycle}{nombre.group(1)}').first()
            if niveau is not None:
                return niveau
        # Niveau SAISI mais introuvable : on refuse. Le déduire du semestre
        # reviendrait à corriger silencieusement une donnée du fichier client
        # (« L9 » deviendrait « L1 »), donc à inventer une inscription.
        return None
    cle_semestre = normaliser_cle(semestre_texte).upper().replace('_', '')
    if cle_semestre in _SEMESTRE_VERS_NIVEAU:
        return Niveau.objects.filter(code__iexact=_SEMESTRE_VERS_NIVEAU[cle_semestre]).first()
    return None


def _trouver_semestre(valeur, niveau):
    """Résout un semestre du niveau (``S3``, ``3``, ``Semestre 3``)."""
    from scolarite.models import Semestre

    if niveau is None:
        return None
    texte = normaliser_cle(valeur).upper().replace('_', '')
    nombre = re.search(r'(\d{1,2})', texte or '')
    if nombre:
        numero = int(nombre.group(1))
        semestre = Semestre.objects.filter(niveau=niveau, numero=numero).first()
        if semestre is not None:
            return semestre
    return None


def _trouver_parcours(valeur, formation):
    from scolarite.models import Parcours

    texte = str(valeur or '').strip()
    if not texte:
        return None
    qs = Parcours.objects.filter(ref_formation=formation)
    trouve = qs.filter(intitule__iexact=texte).first() or qs.filter(code__iexact=texte).first()
    if trouve is not None:
        return trouve
    cible = normaliser_cle(texte)
    for parcours in qs:
        if normaliser_cle(parcours.intitule) == cible or normaliser_cle(parcours.code) == cible:
            return parcours
    return None


def _trouver_ou_creer_vague(valeur, appliquer):
    from formations.models import RefVague

    texte = str(valeur or '').strip()
    if not texte:
        return None
    existante = RefVague.objects.filter(libelle__iexact=texte).first()
    if existante is not None or not appliquer:
        return existante
    return RefVague.objects.create(libelle=texte)


def _trouver_ou_creer_groupe(annee, formation, parcours, niveau, vague, nom, appliquer):
    """Groupe pédagogique de la cohorte — créé si absent **et** si on applique."""
    from scolarite.models import Groupe

    texte = str(nom or '').strip()
    if not texte:
        return None
    existant = Groupe.objects.filter(
        annee_academique=annee, ref_formation=formation, niveau=niveau,
        nom__iexact=texte,
    ).first()
    if existant is not None:
        return existant
    if not appliquer:
        return None
    return Groupe.objects.create(
        annee_academique=annee, ref_formation=formation, parcours=parcours,
        niveau=niveau, vague=vague, nom=texte,
    )


# ── Import ───────────────────────────────────────────────────────────────────

@dataclass
class _LigneResolue:
    ligne: int
    matricule: str
    nom: str
    prenoms: str
    sexe: str
    date_naissance: object
    filiere_code: str
    formation: object
    parcours: object
    niveau: object
    semestre: object
    groupe: object
    vague: object


def _date_naissance(valeur):
    """Convertit une date de naissance texte, sans jamais lever d'exception."""
    from datetime import date, datetime

    texte = str(valeur or '').strip()
    if not texte:
        return None
    for format_ in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%d/%m/%y'):
        try:
            return datetime.strptime(texte, format_).date()
        except ValueError:
            continue
    # Excel peut fournir un nombre de série.
    try:
        serie = float(Decimal(texte))
    except (InvalidOperation, ValueError):
        return None
    if 20000 < serie < 60000:
        from datetime import timedelta
        return date(1899, 12, 30) + timedelta(days=int(serie))
    return None


def _normaliser_sexe(valeur):
    cle = normaliser_cle(valeur)
    if cle in ('m', 'masculin', 'h', 'homme'):
        return 'M'
    if cle in ('f', 'feminin', 'femme'):
        return 'F'
    return ''


def _resoudre_ligne(ligne, correspondance, annee, ligne_numero, appliquer, erreurs):
    """Résout une ligne brute ; ``None`` si elle doit être rejetée."""
    matricule = _valeur(ligne, correspondance, 'matricule')
    nom = _valeur(ligne, correspondance, 'nom')
    if not matricule:
        erreurs.append(LigneErreur(ligne_numero, REJET_MATRICULE_MANQUANT,
                                    'Aucun matricule fourni pour cette ligne.'))
        return None
    if not nom:
        erreurs.append(LigneErreur(ligne_numero, REJET_NOM_MANQUANT,
                                    f'Matricule {matricule} : nom absent.'))
        return None

    formation_texte = _valeur(ligne, correspondance, 'formation')
    filiere_code = _valeur(ligne, correspondance, 'filiere_code')
    formation = _trouver_formation(formation_texte, filiere_code)
    if formation is None:
        # Distinguer « inconnue » de « connue mais hors périmètre INJS ».
        from formations.models import RefFormation

        hors = RefFormation.objects.filter(intitule__iexact=formation_texte).exists()
        motif = REJET_FORMATION_HORS_PERIMETRE if hors else REJET_FORMATION_ABSENTE
        erreurs.append(LigneErreur(
            ligne_numero, motif,
            f'Matricule {matricule} : formation « {formation_texte or filiere_code} »'
            + (' existe mais n’appartient pas au périmètre INJS.'
               if hors else ' introuvable au référentiel INJS.'),
        ))
        return None

    semestre_texte = _valeur(ligne, correspondance, 'semestre')
    niveau = _trouver_niveau(_valeur(ligne, correspondance, 'niveau'), semestre_texte)
    if niveau is None:
        erreurs.append(LigneErreur(
            ligne_numero, REJET_NIVEAU_ABSENT,
            f'Matricule {matricule} : niveau « {_valeur(ligne, correspondance, "niveau")} » introuvable.',
        ))
        return None

    semestre = _trouver_semestre(semestre_texte, niveau)
    if semestre_texte and semestre is None:
        erreurs.append(LigneErreur(
            ligne_numero, REJET_SEMESTRE_ABSENT,
            f'Matricule {matricule} : semestre « {semestre_texte} » introuvable pour {niveau.code}.',
        ))
        return None

    parcours = _trouver_parcours(_valeur(ligne, correspondance, 'parcours'), formation)
    vague = _trouver_ou_creer_vague(_valeur(ligne, correspondance, 'vague'), appliquer)
    groupe = _trouver_ou_creer_groupe(
        annee, formation, parcours, niveau, vague,
        _valeur(ligne, correspondance, 'groupe'), appliquer,
    )

    return _LigneResolue(
        ligne=ligne_numero,
        matricule=matricule,
        nom=nom,
        prenoms=_valeur(ligne, correspondance, 'prenoms'),
        sexe=_normaliser_sexe(_valeur(ligne, correspondance, 'sexe')),
        date_naissance=_date_naissance(_valeur(ligne, correspondance, 'date_naissance')),
        filiere_code=(filiere_code or formation.filiere_code or '').upper(),
        formation=formation,
        parcours=parcours,
        niveau=niveau,
        semestre=semestre,
        groupe=groupe,
        vague=vague,
    )


def importer_inscrits(*, annee_academique_id, nom_fichier, contenu, appliquer=False,
                      acteur=None):
    """Importe un fichier d'inscrits pour une année académique.

    ``appliquer=False`` (défaut) : **analyse seule**, aucune écriture.
    ``appliquer=True`` : écrit les ``InscriptionImportee`` de façon
    transactionnelle et idempotente (clé année + matricule).

    Retourne toujours un :class:`RapportImport` — jamais une exception sur une
    ligne invalide : un fichier client contient toujours des lignes sales.
    """
    from scolarite.models import AnneeAcademique

    rapport = RapportImport(applique=appliquer, fichier=str(nom_fichier or ''))
    annee = AnneeAcademique.objects.filter(pk=annee_academique_id).first()
    if annee is None:
        rapport.motif_echec = 'ANNEE_INTROUVABLE'
        rapport.detail = {'annee_academique_id': annee_academique_id}
        return rapport

    donnees = contenu if isinstance(contenu, (bytes, bytearray)) else str(contenu).encode('utf-8')
    rapport.empreinte = hashlib.sha256(donnees).hexdigest()

    try:
        lignes = lire_fichier(nom_fichier, donnees)
    except ValueError as exc:
        rapport.motif_echec = str(exc)
        return rapport
    except Exception as exc:  # fichier illisible (zip corrompu, etc.)
        rapport.motif_echec = REJET_FORMAT_NON_SUPPORTE
        rapport.detail = {'erreur': str(exc)[:200]}
        return rapport

    if not lignes:
        rapport.motif_echec = REJET_FICHIER_VIDE
        return rapport

    correspondance, ignorees = map_colonnes(list(lignes[0].keys()))
    rapport.colonnes_reconnues = correspondance
    rapport.colonnes_ignorees = ignorees
    manquantes = [c for c in COLONNES_OBLIGATOIRES if c not in correspondance]
    if manquantes:
        rapport.motif_echec = REJET_COLONNES_ABSENTES
        rapport.detail = {'colonnes_manquantes': manquantes}
        rapport.lignes_lues = len(lignes)
        return rapport

    try:
        _executer_import(
            rapport, lignes, correspondance, annee, appliquer,
        )
    except _AnnulerImport:
        # Aucune écriture conservée ; le rapport explique pourquoi.
        rapport.motif_echec = rapport.motif_echec or REJET_FICHIER_VIDE
        rapport.creees = 0
        rapport.mises_a_jour = 0
        rapport.cohortes = []
        rapport.effectif_total = 0

    return rapport


def _executer_import(rapport, lignes, correspondance, annee, appliquer):
    """Corps transactionnel de l'import (isolé pour tester l'annulation)."""
    from edts.models import InscriptionImportee

    with transaction.atomic():
        rapport.lignes_lues = len(lignes)
        resolues = []
        for index, ligne in enumerate(lignes, start=2):
            resolue = _resoudre_ligne(
                ligne, correspondance, annee, index, appliquer, rapport.erreurs,
            )
            if resolue is not None:
                resolues.append(resolue)

        rapport.lignes_valides = len(resolues)
        rapport.lignes_rejetees = len(rapport.erreurs)

        if appliquer:
            from edts.models import InscriptionImportee

            for resolue in resolues:
                _, creee = InscriptionImportee.objects.update_or_create(
                    annee_academique=annee,
                    matricule=resolue.matricule,
                    defaults={
                        'nom': resolue.nom,
                        'prenoms': resolue.prenoms,
                        'sexe': resolue.sexe,
                        'date_naissance': resolue.date_naissance,
                        'filiere_code': resolue.filiere_code,
                        'ref_formation': resolue.formation,
                        'parcours': resolue.parcours,
                        'niveau': resolue.niveau,
                        'semestre': resolue.semestre,
                        'groupe': resolue.groupe,
                        'vague': resolue.vague,
                        'fichier': rapport.fichier[:255],
                        'empreinte': rapport.empreinte,
                        'ligne_source': resolue.ligne,
                    },
                )
                if creee:
                    rapport.creees += 1
                else:
                    rapport.mises_a_jour += 1

            if not rapport.succes:
                # Aucune ligne exploitable : on ne laisse pas de cohorte
                # partielle derrière soi (groupes/vagues créés à vide).
                raise _AnnulerImport()

        rapport.cohortes = _resume_cohortes(resolues)
        rapport.effectif_total = len(resolues)

    return rapport


class _AnnulerImport(Exception):
    """Annule la transaction d'import : aucune écriture partielle ne subsiste."""


def _resume_cohortes(resolues):
    """Agrège les lignes valides par cohorte pédagogique (pour l'écran)."""
    agregat = {}
    for ligne in resolues:
        cle = (
            ligne.formation.pk,
            ligne.parcours.pk if ligne.parcours else None,
            ligne.niveau.pk,
            ligne.semestre.pk if ligne.semestre else None,
            ligne.groupe.pk if ligne.groupe else None,
            ligne.vague.pk if ligne.vague else None,
        )
        if cle not in agregat:
            agregat[cle] = {
                'formation': ligne.formation.intitule,
                'filiere_code': ligne.filiere_code,
                'parcours': ligne.parcours.intitule if ligne.parcours else None,
                'niveau': ligne.niveau.code,
                'semestre': ligne.semestre.numero if ligne.semestre else None,
                'groupe': ligne.groupe.nom if ligne.groupe else None,
                'vague': ligne.vague.libelle if ligne.vague else None,
                'effectif': 0,
            }
        agregat[cle]['effectif'] += 1
    return sorted(agregat.values(), key=lambda c: (
        c['formation'], c['niveau'], c['groupe'] or '', c['vague'] or '',
    ))


def effectifs_par_groupe(annee_academique_id=None):
    """Effectif **réel** issu des imports, par groupe (contrainte de capacité)."""
    from django.db.models import Count

    from edts.models import InscriptionImportee

    qs = InscriptionImportee.objects.filter(
        statut=InscriptionImportee.Statut.INSCRIT, groupe__isnull=False,
    )
    if annee_academique_id is not None:
        qs = qs.filter(annee_academique_id=annee_academique_id)
    return {
        ligne['groupe_id']: ligne['total']
        for ligne in qs.values('groupe_id').annotate(total=Count('pk'))
    }
