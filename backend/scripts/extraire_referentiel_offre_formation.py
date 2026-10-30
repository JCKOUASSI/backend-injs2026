"""Extrait le référentiel officiel INJS depuis le CLASSEUR Excel de référence.

Source primaire
---------------
``dossier_trvx/INJS D'ABIDJAN — RÉFÉRENTIEL EXHAUSTIF DE L'OFFRE DE FORMATION2.xlsx``
(19 feuilles, dont une masquée). C'est le classeur explicitement désigné comme
référentiel officiel de l'offre de formation.

Repli documenté
---------------
Si le classeur est absent de l'arborescence de travail, le script bascule sur
``RÉFÉRENTIEL EXHAUSTIF DE L'OFFRE DE FORMATION.docx``, qui porte le même
référentiel sous forme de tableaux. Le mode utilisé est écrit dans le JSON
produit (``meta.source_type``), de sorte qu'aucune donnée ne puisse être
attribuée à la mauvaise source.

Usage
-----
    python scripts/extraire_referentiel_offre_formation.py
    python scripts/extraire_referentiel_offre_formation.py --source /chemin/ref.xlsx

Le script est **non destructif** : il lit la source et n'écrit que le JSON.
"""

import argparse
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

#: Un semestre LMD vaut 30 crédits (règle institutionnelle confirmée).
CREDITS_PAR_SEMESTRE = 30

#: Codes de filière validés par la source métier du dépôt
#: (``formations/referentiel_injs_l6.py``). Le classeur ne publie pas de code
#: exploitable pour Andragogie, Loisir et Gérontologie (« — »).
FILIERES_PAR_INTITULE = {
    'ÉDUCATION ET MOTRICITÉ': ('F01', 'EM'),
    'ACTIVITÉS PHYSIQUES ADAPTÉES': ('F02', 'APA'),
    'ENTRAÎNEMENT SPORTIF': ('F03', 'ES'),
    'MANAGEMENT DU SPORT': ('F04', 'MS'),
    'ANDRAGOGIE': ('F05', 'AND'),
    'LOISIR': ('F06', 'LOI'),
    'ENTREPRENEURIAT JEUNESSE ET CONDUITE DE PROJETS': ('F07', 'ECJP'),
    'GÉRONTOLOGIE': ('F08', 'GER'),
}

#: Libellé de semestre du classeur → numéro global S1..S10.
SEMESTRES = {
    'Semestre 1': 1, 'Semestre 2': 2, 'Semestre 3': 3, 'Semestre 4': 4,
    'Semestre 5': 5, 'Semestre 6': 6,
    'Semestre 1 (M1/S7)': 7, 'Semestre 2 (M1/S8)': 8,
    'Semestre 3 (M2/S9)': 9, 'Semestre 4 (M2/S10)': 10,
}

#: Palier → niveau LMD.
NIVEAU_PAR_ANNEE = {'L1': 'L1', 'L2': 'L2', 'L3': 'L3', 'M1': 'M1', 'M2': 'M2'}


def _texte(valeur):
    if valeur is None:
        return ''
    return ' '.join(str(valeur).split())


def _slug(texte):
    """Code technique stable dérivé d'un libellé (jamais présenté comme officiel)."""
    sans_accent = ''.join(
        c for c in unicodedata.normalize('NFD', texte)
        if unicodedata.category(c) != 'Mn'
    )
    return re.sub(r'[^A-Za-z0-9]+', '-', sans_accent).strip('-').upper()


def _decouper_ecue(contenu, nb_attendu=None):
    """Découpe les ECUE d'une UE.

    Le classeur fournit le nombre d'ECUE de chaque UE (colonne « Nb ECUE »).
    Ce nombre est **autoritaire** : il tranche les cas où un intitulé contient
    lui-même un point-virgule (ex. « Judo : Fondamentaux ; Lutte :
    Fondamentaux », comptée pour UNE ECUE).
    """
    contenu = _texte(contenu)
    if not contenu:
        return []
    if re.search(r'ECUE\s*1', contenu):
        morceaux = [m.strip(' ;') for m in re.split(r'\s*ECUE\s*\d\s*:\s*', contenu)]
        morceaux = [m for m in morceaux if m]
        if nb_attendu is None or len(morceaux) == nb_attendu:
            return morceaux
    if nb_attendu == 1:
        return [contenu]
    if ';' in contenu and nb_attendu is None:
        return [m.strip() for m in contenu.split(';') if m.strip()]
    return [contenu]


# ─────────────────────────── Lecture du CLASSEUR ────────────────────────────
def _extraire_xlsx(chemin):
    import openpyxl

    classeur = openpyxl.load_workbook(str(chemin), data_only=True)
    anomalies = []

    def lignes(nom, depuis=1):
        if nom not in classeur.sheetnames:
            anomalies.append({'type': 'FEUILLE_ABSENTE', 'detail': nom})
            return
        feuille = classeur[nom]
        for ligne in feuille.iter_rows(min_row=depuis, values_only=True):
            valeurs = [_texte(c) for c in ligne]
            if any(valeurs):
                yield valeurs

    # ── Feuille 1 : filières, mentions, concours ───────────────────────────
    mentions, filieres, concours = {}, {}, []
    section = ''
    for v in lignes('1. Filières & Concours'):
        texte = ' '.join(x for x in v if x).strip()
        if texte.startswith('1.1.'):
            section = 'mentions'
        elif texte.startswith('1.2.'):
            section = 'filieres'
        elif texte.startswith('1.3.'):
            section = 'concours'
        elif section == 'mentions' and v[0].isdigit():
            mentions[v[2]] = {
                'sigle': v[2], 'libelle': v[1], 'ecole': v[3], 'domaine': v[4],
            }
        elif section == 'filieres' and v[0].isdigit():
            libelle = v[2]
            code, sigle = FILIERES_PAR_INTITULE.get(_texte(libelle).upper(), ('', ''))
            if not code:
                anomalies.append({
                    'type': 'FILIERE_NON_RATTACHEE',
                    'detail': f'Filière « {libelle} » absente de la nomenclature F01-F08.',
                })
                continue
            sigle_source = v[3]
            filieres[code] = {
                'code': code, 'rang': int(v[0]), 'libelle': libelle, 'sigle': sigle,
                'sigle_source': '' if sigle_source in ('—', '-') else sigle_source,
                'ecole': v[1],
                'mention': v[1].split('(')[-1].rstrip(')') if '(' in v[1] else '',
                'diplomes': v[4],
            }
        elif section == 'concours' and v[0] and v[0] != 'Sigle':
            concours.append({
                'sigle': v[0], 'concours': v[1], 'conditions': v[2],
                'age': v[3], 'duree_grade': v[4],
            })

    # ── Feuille 2 : cycle professionnel ────────────────────────────────────
    cycles_pro = []
    for v in lignes('2. Cycle professionnel'):
        if v[0] in ('ENSEPS', 'ENSEP') and v[1]:
            cycles_pro.append({
                'ecole': v[0], 'cycle': v[1], 'condition': v[2], 'duree': v[3],
            })

    # ── Feuille 3 : formations à la carte ──────────────────────────────────
    formations_carte = []
    for v in lignes('3. Formations à la carte'):
        if v[0] in ('ENSEPS', 'ENSEP') and v[1]:
            formations_carte.append({
                'ecole': v[0], 'formation': v[1], 'categorie': v[2],
            })

    # ── Feuille 4 : paliers, semestres, grades ─────────────────────────────
    paliers, grades = [], []
    section = ''
    for v in lignes('4. Semestres & Grades'):
        texte = ' '.join(x for x in v if x).strip()
        if texte.startswith('7.'):
            section = 'grades'
            continue
        if v[0] in ('Licence 1', 'Licence 2', 'Licence 3', 'Master 1', 'Master 2',
                    'Doctorat 1-3', 'Cycle professionnel', 'Professorat de collège (concours)'):
            paliers.append({'palier': v[0], 'semestres': v[1],
                            'credits_an': v[2], 'sanction': v[3]})
        elif section == 'grades' and v[0]:
            grades.append({'grade': v[0], 'niveau': v[1],
                           'credits': v[2], 'acces': v[3]})

    # ── Feuille 5 : organisation ───────────────────────────────────────────
    organisation = []
    for v in lignes('5. Organisation'):
        if len(v) > 1 and v[0] and v[0] != 'Sigle':
            if v[0].startswith('STRUCTURE'):
                continue
            organisation.append({'sigle': v[0], 'intitule': v[1]})

    # ── Feuille 8 : base normalisée des enseignements ──────────────────────
    maquettes, cles, familles = {}, {}, {}
    ents = None
    for v in lignes('8. Base UE (source TCD)'):
        if v and v[0] == 'Filière':
            ents = v
            continue
        if ents is None or not v[0]:
            continue
        filiere, _nom, ecole, palier, annee, semestre = v[0], v[1], v[2], v[3], v[4], v[5]
        ue, contenu, nb_ecue, credits, famille, cle = v[6], v[7], v[8], v[9], v[10], v[11]
        if semestre not in SEMESTRES:
            anomalies.append({
                'type': 'SEMESTRE_INCONNU',
                'detail': f'« {semestre} » pour l’UE « {ue} » — ligne ignorée.',
            })
            continue
        numero = SEMESTRES[semestre]
        try:
            nb = int(nb_ecue) if _texte(nb_ecue) else None
        except ValueError:
            nb = None
        try:
            credit = int(float(credits)) if _texte(credits) else None
        except ValueError:
            credit = None
        ecues = _decouper_ecue(contenu, nb)
        if nb is not None and len(ecues) != nb:
            anomalies.append({
                'type': 'ECUE_INCOHERENTE',
                'detail': (
                    f'{filiere} {semestre} « {ue} » : {len(ecues)} ECUE découpées '
                    f'pour {nb} annoncées.'
                ),
            })
        maquettes.setdefault(filiere, {}).setdefault(str(numero), []).append({
            'ue': ue,
            'ecues': ecues,
            'credits': credit,
            'famille': famille,
            'cle_source': cle,
            'palier': palier,
            'annee': annee,
            'source': f'feuille 8, clé « {cle} »',
        })
        if famille:
            familles.setdefault(_slug(famille), {
                'code': _slug(famille), 'libelle': famille,
            })

    filiere_du_sigle = {}
    for code, filiere in filieres.items():
        if filiere['sigle'] in maquettes:
            filiere_du_sigle[filiere['sigle']] = code

    # ── Feuille masquée « Listes » : valeurs de référence ──────────────────
    listes = {}
    entetes = None
    for v in lignes('Listes'):
        if v[0] == 'Cursus':
            entetes = v
            for nom in entetes:
                if nom:
                    listes.setdefault(nom, [])
            continue
        if entetes is None:
            continue
        for i, nom in enumerate(entetes):
            if nom and i < len(v) and v[i]:
                if v[i] not in listes[nom]:
                    listes[nom].append(v[i])

    # ── Feuille 13 : plages horaires du modèle d'emploi du temps ───────────
    plages = []
    for v in lignes('13. Emploi du temps (modèle)'):
        m = re.match(r'^(\d{1,2})[h:](\d{2})\s*[-–]\s*(\d{1,2})[h:](\d{2})$', v[0])
        if m:
            plages.append({
                'libelle': v[0],
                'debut': f'{int(m.group(1)):02d}:{m.group(2)}',
                'fin': f'{int(m.group(3)):02d}:{m.group(4)}',
            })

    # ── Feuille 11 : pondérations provisoires du classeur ──────────────────
    # Le classeur écrit lui-même « à ajuster selon le règlement des études » :
    # ces valeurs sont donc PROVISOIRES et ne doivent jamais être appliquées
    # sans décision institutionnelle.
    notation_provisoire = {
        'statut': 'A_CONFIRMER',
        'source': 'feuille 11 « Suivi des notes », cellule F2/G2',
        'avertissement_source': (
            'Seuil de validation = 10/20 (règle LMD classique — à confirmer avec '
            'le règlement intérieur de l’INJS), feuille 12.'
        ),
        'valeurs': {},
    }
    for v in lignes('11. Suivi des notes'):
        for i, cellule in enumerate(v):
            etiquette = _texte(cellule).lower()
            if 'poids contrôle continu' in etiquette or 'poids controle continu' in etiquette:
                valeur = next((x for x in v[i + 1:i + 3] if _texte(x)), '')
                notation_provisoire['valeurs']['poids_controle_continu'] = valeur
            elif 'poids examen' in etiquette:
                valeur = next((x for x in v[i + 1:i + 3] if _texte(x)), '')
                notation_provisoire['valeurs']['poids_examen'] = valeur
            elif 'ue validées' in etiquette or 'ue validees' in etiquette:
                notation_provisoire['valeurs']['seuil_validation_ue'] = (
                    _texte(cellule).split('>=')[-1].strip(' )')
                )
        if notation_provisoire['valeurs']:
            break

    # Feuille 12 : seuil de validation, formulé « UE validées (>= 10) ».
    import re as _re

    for v in lignes('12. Relevé de notes'):
        texte = ' '.join(v)
        correspondance = _re.search(r'UE validées\s*\(>=\s*(\d+(?:[.,]\d+)?)\s*\)', texte)
        if correspondance:
            notation_provisoire['valeurs']['seuil_validation_ue'] = (
                correspondance.group(1).replace(',', '.')
            )
            break

    # ── Contrôle : 30 crédits par couple filière × semestre ────────────────
    for filiere, par_semestre in maquettes.items():
        for numero in range(1, 11):
            lignes_semestre = par_semestre.get(str(numero))
            if lignes_semestre is None:
                anomalies.append({
                    'type': 'SEMESTRE_ABSENT', 'filiere': filiere, 'semestre': numero,
                    'detail': (
                        f'Le référentiel ne contient aucune UE pour {filiere} S{numero}. '
                        'Aucune UE inventée.'
                    ),
                })
                continue
            total = sum(ligne['credits'] or 0 for ligne in lignes_semestre)
            if total != CREDITS_PAR_SEMESTRE:
                anomalies.append({
                    'type': 'TOTAL_CREDITS', 'filiere': filiere, 'semestre': numero,
                    'detail': f'Somme = {total} au lieu de {CREDITS_PAR_SEMESTRE}.',
                })

    return {
        'meta': {
            'titre': "Référentiel exhaustif de l'offre de formation — INJS Abidjan",
            'source_fichier': chemin.name,
            'source_type': 'xlsx',
            'source_feuilles': classeur.sheetnames,
            'source_feuilles_masquees': [
                f.title for f in classeur.worksheets if f.sheet_state != 'visible'
            ],
            'extrait_le': date.today().isoformat(),
            'extracteur': 'scripts/extraire_referentiel_offre_formation.py',
            'note_codes': (
                'Le référentiel ne publie aucun code officiel d’UE ni d’ECUE. Les codes '
                'du jeu de données sont TECHNIQUES, dérivés du rang : ils ne doivent pas '
                'être présentés comme officiels. La clé métier du classeur '
                '« Filière|Semestre|UE » est conservée telle quelle (cle_source).'
            ),
            'note_volumes': (
                'Les volumes CM/TD/TP ne sont pas publiés : ils restent à 0. '
                'Les pondérations CC/Examen du classeur (0,4 / 0,6) sont signalées '
                '« à confirmer par le règlement des études » et ne sont pas appliquées.'
            ),
        },
        'mentions': mentions,
        'filieres': filieres,
        'concours': concours,
        'cycles_professionnels': cycles_pro,
        'formations_a_la_carte': formations_carte,
        'paliers': paliers,
        'grades': grades,
        'organisation': organisation,
        'familles_modules': familles,
        'listes': listes,
        'plages_horaires': plages,
        'notation_provisoire': notation_provisoire,
        'maquettes': maquettes,
        'filiere_du_sigle': filiere_du_sigle,
        'anomalies': anomalies,
    }


# ───────────────────── Repli : lecture du document Word ─────────────────────
def _extraire_docx(chemin):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(chemin))

    def blocs():
        for enfant in document.element.body.iterchildren():
            if enfant.tag == qn('w:p'):
                yield Paragraph(enfant, document)
            elif enfant.tag == qn('w:tbl'):
                yield Table(enfant, document)

    def cellules(ligne):
        valeurs, precedente = [], object()
        for cellule in ligne.cells:
            texte = ' '.join(cellule.text.split())
            if texte != precedente:
                valeurs.append(texte)
            precedente = texte
        return valeurs

    mentions, filieres, maquettes, anomalies = {}, {}, {}, {}
    anomalies = []
    titre1 = titre2 = ''
    numero_tableau = 0
    for bloc in blocs():
        if isinstance(bloc, Paragraph):
            texte = bloc.text.strip()
            if texte and bloc.style.name == 'Heading 1':
                titre1, titre2 = texte, ''
            elif texte and bloc.style.name == 'Heading 2':
                titre2 = texte
            continue
        numero_tableau += 1
        if numero_tableau == 1:
            for i, ligne in enumerate(bloc.rows):
                c = cellules(ligne)
                if i == 0 or len(c) < 5:
                    continue
                mentions[c[2]] = {'sigle': c[2], 'libelle': c[1],
                                  'ecole': c[3], 'domaine': c[4]}
            continue
        if numero_tableau == 2:
            for i, ligne in enumerate(bloc.rows):
                c = cellules(ligne)
                if i == 0 or len(c) < 5:
                    continue
                code, sigle = FILIERES_PAR_INTITULE.get(_texte(c[2]).upper(), ('', ''))
                if not code:
                    anomalies.append({'type': 'FILIERE_NON_RATTACHEE',
                                      'detail': f'Filière « {c[2]} » non rattachée.'})
                    continue
                filieres[code] = {
                    'code': code, 'rang': int(c[0]), 'libelle': c[2], 'sigle': sigle,
                    'sigle_source': '' if c[3] in ('—', '-') else c[3],
                    'ecole': c[1],
                    'mention': c[1].split('(')[-1].rstrip(')') if '(' in c[1] else '',
                    'diplomes': c[4],
                }
            continue
        entete = re.match(r'11\.(\w+)\s*—', titre1)
        if not entete or titre2 not in SEMESTRES:
            continue
        sigle = entete.group(1)
        numero = SEMESTRES[titre2]
        lignes = []
        for i, ligne in enumerate(bloc.rows):
            c = cellules(ligne)
            if i == 0 and 'UE' in c[0].upper() and 'Crédit' in c[-1]:
                continue
            if len(c) < 3:
                continue
            chiffres = re.search(r'\d+', c[2])
            lignes.append({
                'ue': c[0], 'ecues': _decouper_ecue(c[1]),
                'credits': int(chiffres.group()) if chiffres else None,
                'famille': '', 'cle_source': '', 'palier': '', 'annee': '',
                'source': f'tableau {numero_tableau} ({titre1} / {titre2})',
            })
        maquettes.setdefault(sigle, {})[str(numero)] = lignes

    for sigle, par_semestre in maquettes.items():
        for numero in range(1, 11):
            if str(numero) not in par_semestre:
                anomalies.append({
                    'type': 'SEMESTRE_ABSENT', 'filiere': sigle, 'semestre': numero,
                    'detail': f'Aucune maquette {sigle} S{numero}. Aucune UE inventée.',
                })
                continue
            total = sum(l['credits'] or 0 for l in par_semestre[str(numero)])
            if total != CREDITS_PAR_SEMESTRE:
                anomalies.append({
                    'type': 'TOTAL_CREDITS', 'filiere': sigle, 'semestre': numero,
                    'detail': f'Somme = {total} au lieu de {CREDITS_PAR_SEMESTRE}.',
                })

    return {
        'meta': {
            'titre': "Référentiel exhaustif de l'offre de formation — INJS Abidjan",
            'source_fichier': chemin.name,
            'source_type': 'docx',
            'extrait_le': date.today().isoformat(),
            'extracteur': 'scripts/extraire_referentiel_offre_formation.py',
            'note_codes': 'Codes d’UE/ECUE techniques, non officiels (repli .docx).',
            'note_volumes': 'Volumes CM/TD/TP non publiés : laissés à 0.',
        },
        'mentions': mentions, 'filieres': filieres, 'concours': [],
        'cycles_professionnels': [], 'formations_a_la_carte': [],
        'paliers': [], 'grades': [], 'organisation': [],
        'familles_modules': {}, 'listes': {}, 'plages_horaires': [],
        'notation_provisoire': {'statut': 'INDISPONIBLE', 'valeurs': {}},
        'maquettes': maquettes,
        'filiere_du_sigle': {f['sigle']: c for c, f in filieres.items()},
        'anomalies': anomalies,
    }


def extraire(chemin_source):
    chemin_source = Path(chemin_source)
    if chemin_source.suffix.lower() in ('.xlsx', '.xlsm'):
        return _extraire_xlsx(chemin_source)
    return _extraire_docx(chemin_source)


def main(argv=None):
    racine = Path(__file__).resolve().parents[1]
    dossier = racine.parent / 'dossier_trvx'
    analyseur = argparse.ArgumentParser(description=__doc__)
    analyseur.add_argument(
        '--source', default=None,
        help='Fichier source (.xlsx prioritaire, repli .docx).',
    )
    analyseur.add_argument(
        '--sortie',
        default=str(racine / 'scolarite' / 'data' / 'referentiel_offre_formation.json'),
    )
    options = analyseur.parse_args(argv)

    if options.source:
        candidats = [Path(options.source)]
    else:
        candidats = [
            dossier / "INJS D'ABIDJAN — RÉFÉRENTIEL EXHAUSTIF DE L'OFFRE DE FORMATION2.xlsx",
            dossier / "RÉFÉRENTIEL EXHAUSTIF DE L'OFFRE DE FORMATION.docx",
        ]

    source = next((c for c in candidats if c.is_file()), None)
    if source is None:
        print('Aucune source trouvée :', ', '.join(str(c) for c in candidats), file=sys.stderr)
        return 2

    donnees = extraire(source)
    sortie = Path(options.sortie)
    sortie.parent.mkdir(parents=True, exist_ok=True)
    sortie.write_text(json.dumps(donnees, ensure_ascii=False, indent=1) + '\n',
                      encoding='utf-8')

    total_ue = sum(len(l) for s in donnees['maquettes'].values() for l in s.values())
    total_ecue = sum(len(l['ecues']) for s in donnees['maquettes'].values()
                     for l in s.values() for l in l) if False else sum(
        len(ligne['ecues'])
        for par_semestre in donnees['maquettes'].values()
        for lignes in par_semestre.values()
        for ligne in lignes
    )
    total_credits = sum(ligne['credits'] or 0
                        for par_semestre in donnees['maquettes'].values()
                        for lignes in par_semestre.values()
                        for ligne in lignes)

    print(f'Source            : {source.name} ({donnees["meta"]["source_type"]})')
    if donnees['meta'].get('source_feuilles'):
        print(f'Feuilles          : {len(donnees["meta"]["source_feuilles"])} '
              f'(masquées : {", ".join(donnees["meta"]["source_feuilles_masquees"]) or "aucune"})')
    print(f'Filières          : {len(donnees["filieres"])}')
    print(f'Maquettes         : {sum(len(s) for s in donnees["maquettes"].values())} (filière × semestre)')
    print(f'UE                : {total_ue}')
    print(f'ECUE              : {total_ecue}')
    print(f'Crédits           : {total_credits}')
    print(f'Familles modules  : {len(donnees["familles_modules"])}')
    print(f'Anomalies         : {len(donnees["anomalies"])}')
    for anomalie in donnees['anomalies']:
        print(f'  - [{anomalie["type"]}] {anomalie["detail"]}')
    print(f'Écrit             : {sortie}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
